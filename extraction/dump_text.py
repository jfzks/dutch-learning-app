"""Dump every .docx and .pdf in the source folder to plain text files in extraction/raw/.
We'll then read those .txt files to build the structured JSON content."""
import os, sys, re
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "extraction" / "raw"
OUT.mkdir(parents=True, exist_ok=True)

import fitz  # pymupdf
import docx

def safe_name(p: Path) -> str:
    rel = p.relative_to(ROOT)
    return str(rel).replace("/", "__").replace(" ", "_") + ".txt"

def dump_pdf(p: Path):
    text_parts = []
    with fitz.open(p) as doc:
        for i, page in enumerate(doc):
            text_parts.append(f"\n===== PAGE {i+1} =====\n")
            text_parts.append(page.get_text())
    return "".join(text_parts)

def dump_docx(p: Path):
    d = docx.Document(p)
    parts = []
    for para in d.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    # tables (wordlists are usually tables)
    for ti, table in enumerate(d.tables):
        parts.append(f"\n===== TABLE {ti+1} =====")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)

skipped = []
for p in ROOT.rglob("*"):
    if p.is_dir(): continue
    if p.name.startswith(".") or p.name.startswith("~$"): continue
    if "/app/" in str(p) or "/extraction/" in str(p): continue
    suffix = p.suffix.lower()
    try:
        if suffix == ".pdf":
            text = dump_pdf(p)
        elif suffix == ".docx":
            text = dump_docx(p)
        else:
            continue
        (OUT / safe_name(p)).write_text(text, encoding="utf-8")
    except Exception as e:
        skipped.append((str(p), str(e)))

print(f"Dumped to {OUT}")
if skipped:
    print("Skipped:")
    for s, e in skipped: print(f"  {s}: {e}")
